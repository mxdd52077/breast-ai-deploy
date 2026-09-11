"""Offline characterization probes of the unmodified supplied snapshot.

These assert that reported defects exist; PASS does NOT mean safe behavior.
Run from workspace: .venv/Scripts/python.exe audit/reproduce_findings.py
"""
import ast
import json
import sys
from datetime import date
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "source/breast_roi_copilot"
sys.path.insert(0, str(SRC))

from docx import Document
from pydantic import ValidationError
from src.patient_companion.documents import parse_document
from src.patient_companion.provider import assemble_patient_case
from src.patient_companion.report import build_report
from src.patient_companion.schemas import (
    CareEvent, ClinicalFact, ConfirmationStatus as Status,
    PatientExtraction, PatientProfile,
)
from src.patient_companion.tasks import build_care_tasks

results = []

def probe(name, fn):
    try:
        detail = fn()
        results.append({"id": name, "reproduced": True, "detail": detail})
    except Exception as exc:
        results.append({"id": name, "reproduced": False, "error": repr(exc)})

def fixture():
    doc = parse_document("fictional.txt", "2026年9月12日复诊。".encode())
    fact = ClinicalFact(id="f1", category="复诊", value="2026年9月12日复诊",
        original_text="2026年9月12日复诊。", source_document_id=doc.id, source_page=1)
    profile = PatientProfile(display_name="虚构患者", treatment_stage="未知", diagnosis_summary="未经核对的模型摘要")
    return doc, fact, profile

def status_bypass():
    doc, fact, profile = fixture()
    fact.confirmation_status = Status.CONFIRMED
    event = CareEvent(id="e1", event_type="复诊", title="复诊", scheduled_date=date(2026,9,12), source_fact_ids=["f1"], confirmation_status=Status.CONFIRMED)
    case = assemble_patient_case([doc], PatientExtraction(profile=profile,facts=[fact],events=[event]))
    assert len(build_care_tasks(case.events)) == 1
    return "Model-owned confirmed status survives assembly and produces a task without human review."

def semantic_mismatch():
    doc, fact, profile = fixture()
    fact.value = "2027年1月1日复诊"
    event = CareEvent(id="e1",event_type="复诊",title="另一项没有来源的安排",scheduled_date=date(2027,1,1),source_fact_ids=["f1"])
    case = assemble_patient_case([doc],PatientExtraction(profile=profile,facts=[fact],events=[event]))
    assert case.events[0].scheduled_date.year == 2027
    return "Existing quote permits inconsistent normalized fact/event/date."

def report_bypass():
    doc, fact, profile = fixture()
    case = assemble_patient_case([doc],PatientExtraction(profile=profile,facts=[fact],events=[]))
    report = build_report(case)
    assert not report.source_fact_ids and report.overview == profile.diagnosis_summary
    return "Report overview includes unconfirmed, unreferenced profile summary."

def duplicate_ids():
    doc, fact, profile = fixture()
    case = assemble_patient_case([doc],PatientExtraction(profile=profile,facts=[fact,fact.model_copy()],events=[]))
    assert len(case.facts) == 2 and case.facts[0].id == case.facts[1].id
    return "Duplicate fact IDs accepted; UI widget keys and review mapping cannot disambiguate them."

def docx_table():
    stream = BytesIO()
    doc = Document()
    doc.add_paragraph("虚构记录")
    doc.add_table(rows=1,cols=1).cell(0,0).text = "2026年9月12日复诊"
    doc.save(stream)
    parsed = parse_document("fictional.docx",stream.getvalue())
    assert "复诊" not in parsed.pages[0].text
    return "DOCX table contents silently omitted."

def rename_duplicates():
    content = "虚构记录".encode()
    assert parse_document("a.txt",content).id != parse_document("b.txt",content).id
    return "Same bytes under different filenames receive distinct document IDs."

def missing_date():
    try:
        CareEvent(id="e1",event_type="复诊",title="复诊日期待定",source_fact_ids=["f1"])
    except ValidationError:
        return "Schema cannot represent an event with unknown date; structured extraction must omit it or fail."
    raise AssertionError("Unexpectedly accepts missing date")

def upload_clear():
    tree = ast.parse((SRC/"app_pages/patient_profile.py").read_text(encoding="utf-8"))
    fn = next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=="_clear_case_when_upload_changes")
    state = SimpleNamespace(patient_case="existing",patient_task_status={"task-e1":"已完成"})
    env = {"st":SimpleNamespace(session_state=state)}
    exec(compile(ast.Module(body=[fn],type_ignores=[]),"snapshot callback","exec"),env)
    env[fn.name]()
    assert state.patient_case is None and state.patient_task_status == {}
    return "Actual upload-change callback clears both existing case and task states before replacement succeeds."

def repeated_event():
    e = CareEvent(id="e1",event_type="用药",title="虚构安排",scheduled_date=date(2026,9,1),frequency="每日一次",source_fact_ids=["f1"],confirmation_status=Status.CONFIRMED)
    a = build_care_tasks([e],today=date(2026,9,1))
    b = build_care_tasks([e],today=date(2026,9,5))
    assert a == b and len(a) == 1
    return "Frequency is display text only; today has no effect and recurrence is not expanded."

for name, fn in [("P0-02",status_bypass),("P0-03",semantic_mismatch),("P0-04",report_bypass),
                 ("P1-01",upload_clear),("P1-04",docx_table),("P1-05",missing_date),
                 ("P1-06a",rename_duplicates),("P1-06b",duplicate_ids),("P1-07",repeated_event)]:
    probe(name,fn)

out = ROOT/"deliverables/verification"
out.mkdir(parents=True,exist_ok=True)
(out/"probes.json").write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps(results,ensure_ascii=False,indent=2))
sys.exit(0 if all(r["reproduced"] for r in results) else 1)
