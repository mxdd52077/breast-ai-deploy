from io import BytesIO
import json
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError
from docx import Document as WordDocument
from PIL import Image

from app.backend.main import app
from app.backend.db import Session,Document,DocumentChunk,Fact,Task,Job,DATA
from app.backend.worker import process_one
from app.backend import providers
from app.backend.providers import parse_pages,extract_facts,absolute_date,infer_report_date,ServiceError,ocr_document
from app.backend.contracts import Extraction,Answer

def workspace(c):
    r=c.get('/api/workspace');assert r.status_code==200,r.text
    return r.json()

def review(c,f,status='confirmed'):
    return c.patch('/api/facts/'+f['id'],json={'status':status,'version':f['version'],'note':'已对照虚构测试资料'})

def upload(c,text,name='测试资料.txt'):
    r=c.post('/api/documents',files={'file':(name,text.encode(),'text/plain')})
    assert r.status_code==200,r.text
    return r.json()

def test_isolation_and_csrf(demo):
    w=workspace(demo);doc=w['documents'][0];fact=w['facts'][0]
    with TestClient(app) as other:
        a=other.post('/api/auth/demo').json();other.headers['X-CSRF-Token']=a['csrf']
        assert workspace(other)['user']['id']!=w['user']['id']
        for suffix in ['', '/file']:
            assert other.get('/api/documents/'+doc['id']+suffix).status_code==404
        assert other.patch('/api/facts/'+fact['id'],json={'status':'confirmed','version':1}).status_code==404
    assert demo.patch('/api/facts/'+fact['id'],headers={'X-CSRF-Token':'wrong'},json={'status':'confirmed','version':1}).status_code==403
    assert demo.post('/api/documents/'+doc['id']+'/retry',headers={'Origin':'https://evil.example'}).status_code==403

def test_task_completion_survives_review_upload_and_refresh(demo):
    f=next(f for f in workspace(demo)['facts'] if f['scheduled_date'])
    assert not workspace(demo)['tasks']
    assert review(demo,f).status_code==200
    t=workspace(demo)['tasks'][0]
    assert demo.patch('/api/tasks/'+t['id'],json={'status':'completed','version':t['version']}).status_code==200
    f=next(x for x in workspace(demo)['facts'] if x['id']==f['id'])
    assert review(demo,f).status_code==200
    d=upload(demo,'检查安排：日期尚待明确。');assert process_one()
    assert upload(demo,'检查安排：日期尚待明确。','改名.txt')['duplicate']
    w=workspace(demo)
    assert len(w['documents'])==3
    assert w['tasks'][0]['status']=='completed'
    assert len(w['tasks'])==1
    assert any(x['document_id']==d['id'] and x['scheduled_date'] is None for x in w['facts'])
    assert review(demo,f).status_code==409
    fresh=next(x for x in w['facts'] if x['id']==f['id'])
    assert review(demo,fresh,'pending').status_code==200
    stale_task=w['tasks'][0]
    assert demo.patch('/api/tasks/'+stale_task['id'],json={'status':'pending','version':stale_task['version']}).status_code==409
    latest=next(x for x in workspace(demo)['facts'] if x['id']==f['id'])
    assert review(demo,latest).status_code==200
    assert workspace(demo)['tasks'][0]['status']=='completed'

def test_delete_document_cascades(demo):
    f=next(f for f in workspace(demo)['facts'] if f['scheduled_date'])
    assert review(demo,f).status_code==200
    assert demo.delete('/api/documents/'+f['document_id']).status_code==200
    w=workspace(demo)
    assert not w['tasks']
    assert all(x['document_id']!=f['document_id'] for x in w['facts'])
    assert demo.get('/api/documents/'+f['document_id']+'/file').status_code==404

def test_account_gate_and_session(client):
    import uuid
    credentials={'username':'test-'+uuid.uuid4().hex,'password':'only-test-password','name':'测试'}
    r=client.post('/api/auth/register',json=credentials);assert r.status_code==200,r.text
    client.headers['X-CSRF-Token']=r.json()['csrf']
    assert client.get('/api/institution/defaults').status_code==403
    assert client.post('/api/documents',files={'file':('x.txt',b'test')}).status_code==503
    assert client.post('/api/auth/logout').status_code==200
    assert client.get('/api/workspace').status_code==401
    r=client.post('/api/auth/login',json=credentials);assert r.status_code==200
    client.headers['X-CSRF-Token']=r.json()['csrf']
    assert client.delete('/api/account').status_code==200
    assert client.get('/api/workspace').status_code==401
    assert client.post('/api/auth/login',json=credentials).status_code==401

def test_missing_ocr_retry_recovery(demo,monkeypatch):
    b=BytesIO();Image.new('RGB',(20,20),'white').save(b,format='PNG')
    r=demo.post('/api/documents',files={'file':('scan.png',b.getvalue(),'image/png')});key=r.json()['id']
    assert process_one()
    d=next(d for d in workspace(demo)['documents'] if d['id']==key)
    assert d['status']=='failed' and '尚未连接' in d['error']
    assert demo.get('/api/documents/'+key+'/file').status_code==200
    monkeypatch.setattr('app.backend.providers.ocr_document',lambda p:[{'page':1,'text':'复诊安排：2030年9月10日上午9点复诊。','location':'第1页 · OCR识别','ocr':True}])
    assert demo.post('/api/documents/'+key+'/retry').status_code==200
    assert process_one()
    w=workspace(demo);assert next(d for d in w['documents'] if d['id']==key)['status']=='ready'
    assert all(f['status']=='pending' for f in w['facts'] if f['document_id']==key)

def test_vercel_can_resume_queued_document(demo,monkeypatch):
    created=upload(demo,'复诊安排：日期尚待明确。','等待整理.txt')
    called=[]
    monkeypatch.setenv('VERCEL','1')
    monkeypatch.setattr('app.backend.main.process_one',lambda key: called.append(key))
    assert demo.post('/api/documents/'+created['id']+'/retry').status_code==200
    assert called==[created['id']]

def test_docx_table_and_unknown_dates(tmp_path):
    d=WordDocument();d.add_paragraph('虚构测试资料');t=d.add_table(rows=1,cols=2);t.cell(0,0).text='复诊';t.cell(0,1).text='2030年9月10日复查'
    p=tmp_path/'table.docx';d.save(p)
    pages=parse_pages(p,'.docx');assert '2030年9月10日' in pages[0]['text']
    facts,_=extract_facts(pages,demo=True);assert any(f['scheduled_date']=='2030-09-10' for f in facts)
    assert absolute_date('下周复诊')==(None,None)
    assert absolute_date('9月10日复诊')==(None,None)
    assert absolute_date('2030年2月30日复诊')==(None,None)
    assert absolute_date('2030年9月10日或2030年9月11日复诊')==(None,None)

def test_extraction_cannot_confirm_or_invent_quote(monkeypatch):
    with pytest.raises(ValidationError):
        Extraction.model_validate({'facts':[{'category':'复诊','page':1,'quote':'test','status':'confirmed'}]})
    monkeypatch.setattr('app.backend.providers.chat_json',lambda *a:Extraction.model_validate({'facts':[{'category':'复诊','page':1,'quote':'不存在的安排'}]}))
    with pytest.raises(ServiceError): extract_facts([{'page':1,'text':'实际原文','location':'第1页','ocr':False}])

def test_typed_ocr_adapter(monkeypatch,tmp_path):
    monkeypatch.setenv('PADDLEOCR_ACCESS_TOKEN','fake-for-adapter-only')
    fake=SimpleNamespace(parse_document=lambda **kw:SimpleNamespace(pages=[SimpleNamespace(markdown_text='测试原文')]),close=lambda:None)
    with patch('paddleocr.PaddleOCRClient',return_value=fake) as constructor:
        pages=ocr_document(tmp_path/'fake.png')
        assert pages[0]['text']=='测试原文'
        assert constructor.call_args.kwargs['base_url']=='https://paddleocr.aistudio-app.com'

def test_grounded_history_and_invalid_citation(demo,monkeypatch):
    r=demo.post('/api/messages',json={'question':'复诊前可以准备哪些问题？'});assert r.status_code==200,r.text
    history=demo.get('/api/messages').json();assert len(history)==2
    assert history[-1]['citations']
    assert any(c['url'].startswith('https://') for c in history[-1]['citations'] if c['source_type']=='knowledge')
    monkeypatch.setenv('APEX_LLM_API_KEY','fake-test-only')
    monkeypatch.setattr('app.backend.main.chat_json',lambda *a:Answer(answer='不受支持的回答',status='supported',citations=['fabricated']))
    assert demo.post('/api/messages',json={'question':'复诊问题准备'}).status_code==503
    assert len(demo.get('/api/messages').json())==2
    assert demo.post('/api/messages',json={'question':'现在呼吸困难'}).status_code==200
    assert demo.get('/api/messages').json()[-1]['status']=='safety_escalation'

def test_chat_ndjson_stream_persists_verified_answer(demo):
    response=demo.post('/api/messages/stream',json={'question':'复诊前准备什么？'})
    assert response.status_code==200
    events=[json.loads(line) for line in response.text.splitlines()]
    assert events[0]['type']=='status'
    assert ''.join(e['text'] for e in events if e['type']=='delta')
    assert events[-1]['type']=='done'
    history=demo.get('/api/messages').json()
    assert history[-2]['role']=='user' and history[-1]['role']=='assistant'

def test_personal_retrieval_requires_confirmation_and_is_account_scoped(demo):
    marker='APEXZETA927'
    created=upload(demo,f'复诊记录：携带{marker}检查单。','个人复诊记录.txt')
    assert process_one(created['id'])
    fact=next(f for f in workspace(demo)['facts'] if f['document_id']==created['id'])
    before=demo.post('/api/messages',json={'question':marker}).json()
    assert before['ok']
    assert not any(c['source_type']=='personal' and c['document_id']==created['id'] for c in demo.get('/api/messages').json()[-1]['citations'])
    assert review(demo,fact).status_code==200
    assert demo.post('/api/messages',json={'question':marker}).status_code==200
    citations=demo.get('/api/messages').json()[-1]['citations']
    personal=next(c for c in citations if c['source_type']=='personal')
    assert personal['document_id']==created['id'] and personal['page']==1
    with TestClient(app) as other:
        session=other.post('/api/auth/demo').json();other.headers['X-CSRF-Token']=session['csrf']
        assert other.post('/api/messages',json={'question':marker}).status_code==200
        assert not any(c['source_type']=='personal' and c['document_id']==created['id'] for c in other.get('/api/messages').json()[-1]['citations'])

def test_roi_matches_original(demo):
    from dataclasses import asdict
    from src.models import BreastROIInputs,calculate_breast_roi
    i=demo.get('/api/institution/defaults').json()
    r=demo.post('/api/institution/simulations',json={'name':'测试方案','inputs':i})
    assert r.status_code==200,r.text
    assert r.json()['results']==asdict(calculate_breast_roi(BreastROIInputs(**i)))
    assert demo.get('/api/institution/simulations').json()[0]['id']==r.json()['id']
    assert demo.post('/api/institution/simulations',json={'inputs':{**i,'target_screening_rate':200}}).status_code==400

def test_batch_confirmation_atomic_and_creates_only_dated_tasks(demo):
    pending=[f for f in workspace(demo)['facts'] if f['status']=='pending']
    items=[{'id':f['id'],'version':f['version']} for f in pending]
    stale=[*items[:-1],{**items[-1],'version':99}]
    assert demo.post('/api/facts/confirm-batch',json={'items':stale}).status_code==409
    assert all(f['status']=='pending' for f in workspace(demo)['facts'] if f['id'] in {i['id'] for i in items})
    assert demo.post('/api/facts/confirm-batch',json={'items':items}).status_code==200
    after=workspace(demo)
    assert all(f['status']=='confirmed' for f in after['facts'])
    assert len(after['tasks'])==sum(bool(f['scheduled_date']) for f in pending)
    assert demo.post('/api/facts/confirm-batch',json={'items':items}).status_code==409
    assert len(workspace(demo)['tasks'])==len(after['tasks'])

def test_kimi_schedule_requires_source_date(monkeypatch):
    monkeypatch.setenv('APEX_LLM_API_KEY','test')
    quote='请于2030年9月10日上午9点完成检查。'
    pages=[{'page':1,'text':quote,'location':'第1页'}]
    def result(day):
        return Extraction.model_validate({'facts':[{'quote':quote,'page':1,'category':'检查','is_schedule':True,'scheduled_date':day,'scheduled_time':'09:00'}]})
    monkeypatch.setattr('app.backend.providers.chat_json',lambda *args:result('2030-09-10'))
    assert extract_facts(pages)[0][0]['scheduled_date']=='2030-09-10'
    monkeypatch.setattr('app.backend.providers.chat_json',lambda *args:result('2030-09-11'))
    assert extract_facts(pages)[0][0]['scheduled_date'] is None

def test_relative_hour_window_uses_later_document_timestamp_as_review_candidate():
    pages=[{'page':1,'location':'第1页 · OCR识别','text':'医嘱：化疗结束后24小时至48小时使用艾多6mg皮下注射一次。医务人员签名：日期：2024年07月10日 09:14。'}]
    infer=getattr(providers,'infer_relative_schedules',lambda _:[])
    item=infer(pages)[0]
    assert item['quote']=='化疗结束后24小时至48小时使用艾多6mg皮下注射一次'
    assert (item['scheduled_date'],item['scheduled_time'])==('2024-07-11','09:14')
    assert (item['scheduled_end_date'],item['scheduled_end_time'])==('2024-07-12','09:14')
    assert '2024-07-10 09:14' in item['schedule_basis']

def test_report_date_prefers_labeled_report_date_and_excludes_patient_dates():
    pages=[{'page':1,'location':'第1页 · OCR识别','text':'出生日期：1977-09-11\n超声检查报告 检查号：240606 日期：2024-06-06\n审核时间：2024-06-06 10:51:21'}]
    assert infer_report_date(pages)=='2024-06-06'
    assert infer_report_date([{'page':1,'location':'第1页','text':'出生日期：1977-09-11；就诊日期：2024-06-05'}]) is None

def test_report_date_is_auto_detected_and_can_be_manually_updated(demo):
    created=upload(demo,'超声检查报告\n日期：2030-09-10\n诊断：测试资料。','带报告日期.txt')
    assert process_one(created['id'])
    document=next(d for d in workspace(demo)['documents'] if d['id']==created['id'])
    assert document['report_date']=='2030-09-10'
    changed=demo.patch('/api/documents/'+created['id']+'/report-date',json={'report_date':'2030-09-12'})
    assert changed.status_code==200,changed.text
    document=next(d for d in workspace(demo)['documents'] if d['id']==created['id'])
    assert document['report_date']=='2030-09-12'
    assert demo.patch('/api/documents/'+created['id']+'/report-date',json={'report_date':'2030-02-30'}).status_code==422
    assert not any(t['fact_id'] in {f['id'] for f in workspace(demo)['facts'] if f['document_id']==created['id']} for t in workspace(demo)['tasks'])

def test_relative_week_schedule_uses_nearest_prior_treatment_date():
    pages=[{'page':1,'location':'第1页 · OCR识别','text':'处理意见：2024-7-9行第1周期PCb-EC方案治疗。离院建议：2、三周后返院行下一周期化疗；日期：2024年07月10日 09:14。'}]
    item=next(x for x in providers.infer_relative_schedules(pages) if x['category']=='治疗')
    assert item['quote']=='三周后返院行下一周期化疗'
    assert (item['scheduled_date'],item['scheduled_time'])==('2024-07-30',None)
    assert item['scheduled_end_date'] is None
    assert '同页前文治疗日期 2024-07-09' in item['schedule_basis']

def test_relative_schedule_merges_with_kimi_quote_prefix(monkeypatch):
    monkeypatch.setenv('APEX_LLM_API_KEY','test')
    text='处理意见：2024-7-9行第1周期方案治疗。离院建议：2、三周后返院行下一周期化疗；'
    quote='2、三周后返院行下一周期化疗'
    pages=[{'page':1,'location':'第1页 · OCR识别','text':text}]
    extracted=Extraction.model_validate({'facts':[{'quote':quote,'page':1,'category':'治疗','is_schedule':True,'scheduled_date':None,'scheduled_time':None}]})
    monkeypatch.setattr('app.backend.providers.chat_json',lambda *args:extracted)
    facts,_=extract_facts(pages)
    treatment=[item for item in facts if '三周后' in item['quote']]
    assert len(treatment)==1
    assert treatment[0]['scheduled_date']=='2024-07-30'

def test_postoperative_and_post_discharge_schedules_use_named_anchors(monkeypatch):
    monkeypatch.setenv('APEX_LLM_API_KEY','test')
    text=('入院后于2024年12月26日在全麻下行左侧乳房重建术，手术顺利。'
          '出院时间：2024年12月27日 07:40。出院医嘱：乳房重建患者术后第五天到华西医院查看切口情况，'
          '出院2周后到门诊病理科领取术后病理报告。')
    quote='乳房重建患者术后第五天到华西医院查看切口情况，出院2周后到门诊病理科领取术后病理报告'
    pages=[{'page':1,'location':'第1页 · OCR识别','text':text}]
    extracted=Extraction.model_validate({'facts':[{'quote':quote,'page':1,'category':'复诊','is_schedule':True,'scheduled_date':None,'scheduled_time':None}]})
    monkeypatch.setattr('app.backend.providers.chat_json',lambda *args:extracted)
    facts,_=extract_facts(pages)
    schedules=[item for item in facts if item['scheduled_date']]
    assert [(item['scheduled_date'],item['scheduled_time']) for item in schedules]==[('2024-12-31',None),('2025-01-10',None)]
    assert all('人工核对' in item['schedule_basis'] for item in schedules)
    assert not any(item['quote']==quote for item in facts)

def test_verified_relative_schedules_survive_unavailable_llm(monkeypatch):
    text=('于2024年12月26日行乳房重建术。出院时间：2024年12月27日。'
          '乳房重建患者术后第五天到医院查看切口，出院2周后到门诊领取病理报告。')
    pages=[{'page':1,'location':'第1页 · OCR识别','text':text}]
    monkeypatch.setattr('app.backend.providers.chat_json',lambda *args:(_ for _ in ()).throw(ServiceError('智能服务未返回可验证结果，资料已保留，可重试。')))
    facts,method=extract_facts(pages)
    assert sorted(item['scheduled_date'] for item in facts)==['2024-12-31','2025-01-10']
    assert method=='时间规则识别 · 待人工核对'

def test_existing_document_can_recognize_relative_schedules_without_reupload(demo):
    text=('于2024年12月26日在全麻下行乳房重建术。出院时间：2024年12月27日。'
          '乳房重建患者术后第五天到医院查看切口，出院2周后到门诊领取病理报告。')
    created=upload(demo,text,'旧资料.txt')
    assert process_one(created['id'])
    with Session() as db:
        facts=db.query(Fact).filter(Fact.document_id==created['id']).all()
        for fact in facts:db.delete(fact)
        db.commit()
    response=demo.post('/api/documents/'+created['id']+'/recognize-schedules')
    assert response.status_code==200,response.text
    assert response.json()['added']==2
    facts=[f for f in workspace(demo)['facts'] if f['document_id']==created['id']]
    assert sorted(f['scheduled_date'] for f in facts)==['2024-12-31','2025-01-10']
    assert all(f['status']=='pending' and f['scheduled_time'] is None for f in facts)
    again=demo.post('/api/documents/'+created['id']+'/recognize-schedules')
    assert again.status_code==200 and again.json()['added']==0

def test_schedule_recognition_ignores_trailing_quote_punctuation(demo):
    text='于2024年12月26日行乳房重建术。乳房重建患者术后第五天到医院查看切口。'
    created=upload(demo,text,'标点测试.txt')
    assert process_one(created['id'])
    first=demo.post('/api/documents/'+created['id']+'/recognize-schedules')
    assert first.status_code==200
    dates=[f['scheduled_date'] for f in workspace(demo)['facts'] if f['document_id']==created['id'] and f['scheduled_date']]
    assert dates==['2024-12-31']

def test_long_ocr_surgery_description_still_anchors_postoperative_followup():
    text=('入院时间：2024年12月25日 07:55。出院时间：2024年12月27日 07:40。'
          '入院后积极完善术前相关准备于2024年12月26日在全麻下行“腔镜下左侧乳腺单纯皮下切除术+'
          '左乳前哨淋巴结切除术活检术+腔镜下胸肌后假体联合补片植入重建术+腔镜下左侧乳头乳晕整复术”。'
          '手术经过顺利。出院医嘱及建议：乳房重建患者术后第五天到医院查看切口情况，'
          '出院2周后到门诊病理科领取术后病理报告。')
    pages=[{'page':1,'location':'第1页 · OCR识别','text':text}]
    items=providers.infer_relative_schedules(pages)
    assert sorted(item['scheduled_date'] for item in items)==['2024-12-31','2025-01-10']
    assert any('手术日期 2024-12-26' in item['schedule_basis'] for item in items)

def test_postoperative_suture_removal_window_is_separate_from_recurring_and_conditional_care():
    text=('于2024年12月26日在全麻下行乳房重建术。出院医嘱：'
          '每周更换伤口敷料2–3次及引流管，提前挂号换药；'
          '重建患者腋窝切口术后10–14天拆线；乳房切口不拆线让其自行脱落或6周后拆线。')
    pages=[{'page':1,'location':'第1页 · OCR识别','text':text}]
    items=providers.infer_relative_schedules(pages)
    assert len(items)==1
    item=items[0]
    assert item['quote']=='术后10–14天拆线'
    assert (item['scheduled_date'],item['scheduled_end_date'])==('2025-01-05','2025-01-09')
    assert item['scheduled_time'] is None and item['scheduled_end_time'] is None
    assert '日期窗口' in item['schedule_basis']

def test_existing_document_backfills_suture_removal_window(demo):
    text=('于2024年12月26日行乳房重建术。'
          '每周更换伤口敷料2-3次；腋窝切口术后10-14天拆线；乳房切口自行脱落或6周后拆线。')
    created=upload(demo,text,'拆线旧资料.txt')
    assert process_one(created['id'])
    with Session() as db:
        for fact in db.query(Fact).filter(Fact.document_id==created['id']).all():db.delete(fact)
        db.commit()
    response=demo.post('/api/documents/'+created['id']+'/recognize-schedules')
    assert response.status_code==200 and response.json()['added']==1
    fact=next(f for f in workspace(demo)['facts'] if f['document_id']==created['id'])
    assert (fact['scheduled_date'],fact['scheduled_end_date'])==('2025-01-05','2025-01-09')
    assert fact['status']=='pending' and '日期窗口' in fact['schedule_basis']

def test_backfill_does_not_duplicate_longer_quote_for_same_window(demo):
    text='于2024年12月26日行乳房重建术。重建患者腋窝切口术后10–14天拆线。'
    created=upload(demo,text,'拆线去重.txt')
    assert process_one(created['id'])
    with Session() as db:
        doc=db.get(Document,created['id'])
        for old in db.query(Fact).filter(Fact.document_id==created['id']).all():db.delete(old)
        db.flush()
        db.add(Fact(user_id=doc.user_id,document_id=doc.id,category='复诊',value='重建患者腋窝切口术后10–14天拆线',quote='重建患者腋窝切口术后10–14天拆线',page=1,status='pending',scheduled_date='2025-01-05',scheduled_end_date='2025-01-09'))
        db.commit()
    result=demo.post('/api/documents/'+created['id']+'/recognize-schedules')
    assert result.status_code==200 and result.json()['added']==0

def test_review_can_set_missing_time_and_window_before_creating_task(demo):
    created=upload(demo,'治疗安排：2030年9月10日进行治疗。','待补时间.txt')
    assert process_one(created['id'])
    fact=next(f for f in workspace(demo)['facts'] if f['document_id']==created['id'])
    payload={'status':'confirmed','version':fact['version'],'note':'已人工核对时间','scheduled_date':'2030-09-10','scheduled_time':'14:30','scheduled_end_date':'2030-09-10','scheduled_end_time':'16:30'}
    response=demo.patch('/api/facts/'+fact['id'],json=payload)
    assert response.status_code==200,response.text
    updated=next(f for f in workspace(demo)['facts'] if f['id']==fact['id'])
    task=next(t for t in workspace(demo)['tasks'] if t['fact_id']==fact['id'])
    assert (updated['scheduled_time'],updated['scheduled_end_time'])==('14:30','16:30')
    assert (task['due_date'],task['due_time'],task['due_end_date'],task['due_end_time'])==('2030-09-10','14:30','2030-09-10','16:30')

def test_review_text_can_be_corrected_without_losing_original_quote(demo):
    fact=next(f for f in workspace(demo)['facts'] if f['scheduled_date'])
    original_quote=fact['quote']
    corrected='人工校正后的乳腺外科复诊安排'
    response=demo.patch('/api/facts/'+fact['id'],json={'status':'confirmed','version':fact['version'],'note':'','value':corrected})
    assert response.status_code==200,response.text
    updated=next(f for f in workspace(demo)['facts'] if f['id']==fact['id'])
    task=next(t for t in workspace(demo)['tasks'] if t['fact_id']==fact['id'])
    assert updated['value']==corrected
    assert updated['quote']==original_quote
    assert task['title']==corrected
    with Session() as db:
        assert db.query(DocumentChunk).filter_by(fact_id=fact['id']).one().content==corrected
    assert demo.patch('/api/facts/'+fact['id'],json={'status':'confirmed','version':updated['version'],'note':'','value':''}).status_code==422

def test_agent_natural_language_calendar_draft_requires_confirmation(demo,monkeypatch):
    class Proposal:
        has_action=True;title='乳腺外科复诊';category='复诊';scheduled_date='2030-09-10';scheduled_time='09:30';scheduled_end_date=None;scheduled_end_time=None;clarification=''
    monkeypatch.setattr('app.backend.main.calendar_proposal',lambda _:Proposal(),raising=False)
    response=demo.post('/api/agent/calendar',json={'question':'请把2030年9月10日上午9点半乳腺外科复诊加入日历'})
    assert response.status_code==200,response.text
    assert response.json()['handled'] is True
    action=demo.get('/api/messages').json()[-1]['actions'][0]
    assert action['status']=='pending' and action['scheduled_time']=='09:30'
    assert not any(task['fact_id']==action['id'] for task in workspace(demo)['tasks'])
    confirmed=demo.patch('/api/facts/'+action['id'],json={'status':'confirmed','version':action['version'],'note':'','scheduled_date':action['scheduled_date'],'scheduled_time':action['scheduled_time'],'scheduled_end_date':None,'scheduled_end_time':None})
    assert confirmed.status_code==200,confirmed.text
    assert any(task['fact_id']==action['id'] for task in workspace(demo)['tasks'])
    assert demo.get('/api/messages').json()[-1]['actions'][0]['status']=='confirmed'

def test_agent_uploaded_document_returns_reviewable_calendar_actions(demo):
    created=upload(demo,'复诊安排：请于2030年9月10日上午9点至乳腺外科复诊。','聊天上传.txt')
    assert process_one(created['id'])
    response=demo.post('/api/agent/documents/'+created['id'],json={'question':'识别图片里的安排并加入日历'})
    assert response.status_code==200,response.text
    payload=response.json()
    assert payload['handled'] is True and payload['action_count']==1
    messages=demo.get('/api/messages').json()
    assert messages[-1]['actions'][0]['scheduled_date']=='2030-09-10'

def test_institution_data_check_aggregates_only(demo):
    content='patient_id,as_of_date,age,last_screen_date,never_screened,has_active_appointment,outreach_consent,is_synthetic\nTEST-PRIVATE-ID,2026-09-01,55,2023-01-01,false,false,true,true\n'
    r=demo.post('/api/institution/data-check',files={'file':('demo.csv',content.encode())})
    assert r.status_code==200,r.text
    assert r.json()['passed'] and r.json()['row_count']==1
    assert 'TEST-PRIVATE-ID' not in r.text
    bad=demo.post('/api/institution/data-check',files={'file':('bad.csv',b'age\n55\n')})
    assert bad.status_code==200 and not bad.json()['passed']
    bad=demo.post('/api/institution/data-check',files={'file':('bad.csv',content.replace(',55,',',90,').encode())})
    assert not bad.json()['passed']

def test_inline_preview_is_authorized_and_typed(demo):
    b=BytesIO();Image.new('RGB',(20,20),'white').save(b,format='PNG')
    key=demo.post('/api/documents',files={'file':('preview.png',b.getvalue())}).json()['id']
    r=demo.get('/api/documents/'+key+'/file?inline=true')
    assert r.status_code==200 and r.headers['content-type']=='image/png'
    assert r.headers['content-disposition'].startswith('inline')
    with TestClient(app) as visitor:
        assert visitor.get('/api/documents/'+key+'/file?inline=true').status_code==401
    # Finish queued fixture so subsequent worker tests do not see it.
    process_one()

def test_docx_embedded_scan_requires_ocr_and_keeps_source(tmp_path,monkeypatch):
    image=BytesIO();Image.new('RGB',(20,20),'white').save(image,format='PNG');image.seek(0)
    doc=WordDocument();doc.add_paragraph('虚构资料正文');doc.add_picture(image)
    path=tmp_path/'embedded.docx';doc.save(path)
    with pytest.raises(ServiceError,match='尚未连接'):parse_pages(path,'.docx')
    monkeypatch.setattr('app.backend.providers.ocr_document',lambda p:[{'page':1,'text':'图片中复诊日期2030年9月10日','location':'第1页','ocr':True}])
    pages=parse_pages(path,'.docx')
    assert len(pages)==2 and pages[0]['text'].strip()=='虚构资料正文'
    assert pages[1]['ocr'] and '嵌入图片1' in pages[1]['location']
    assert pages[1]['text']=='图片中复诊日期2030年9月10日'

def test_nested_docx_tables_are_not_lost(tmp_path):
    doc=WordDocument();table=doc.add_table(rows=1,cols=1)
    nested=table.cell(0,0).add_table(rows=1,cols=1);nested.cell(0,0).text='嵌套表格复诊记录'
    path=tmp_path/'nested.docx';doc.save(path)
    assert '嵌套表格复诊记录' in parse_pages(path,'.docx')[0]['text']
