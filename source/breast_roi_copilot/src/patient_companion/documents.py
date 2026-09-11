"""Deterministic text extraction for supported electronic documents."""

from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from .schemas import DocumentPage, SourceDocument


class DocumentParseError(ValueError):
    """A user-safe document parsing failure."""


MAX_FILE_BYTES = 10 * 1024 * 1024


def _clean_pages(values: list[str]) -> list[DocumentPage]:
    return [
        DocumentPage(page_number=index, text=text.strip())
        for index, text in enumerate(values, start=1)
        if text.strip()
    ]


def parse_document(filename: str, content: bytes) -> SourceDocument:
    if len(content) > MAX_FILE_BYTES:
        raise DocumentParseError("文件超过 10 MB，请压缩或拆分后重试。")
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".txt":
            pages = _clean_pages(content.decode("utf-8-sig").split("\f"))
            media_type = "text/plain"
        elif extension == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(content))
            pages = _clean_pages([(page.extract_text() or "") for page in reader.pages])
            media_type = "application/pdf"
        elif extension == ".docx":
            from docx import Document

            document = Document(BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            pages = _clean_pages([text])
            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            raise DocumentParseError("首版仅支持 PDF、DOCX 和 TXT 标准电子文档。")
    except DocumentParseError:
        raise
    except (UnicodeDecodeError, ValueError, KeyError, OSError) as exc:
        raise DocumentParseError("文件损坏或无法解析，请换用标准电子文档。") from exc

    if not pages:
        raise DocumentParseError("文件没有可读取的文字；扫描件和图片暂不支持。")
    digest = sha256(filename.encode("utf-8") + content).hexdigest()[:16]
    return SourceDocument(
        id=f"doc-{digest}",
        filename=filename,
        media_type=media_type,
        uploaded_at=datetime.now(timezone.utc),
        pages=pages,
    )
