from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/Users/zhangwenqi/Documents/New project/breast_roi_copilot")
OUTPUT = ROOT / "artifacts" / "APEX_AI产品经理面试问答手册_V1.0.docx"

TEAL = "0D7C80"
NAVY = "172033"
LIGHT = "EAF4F4"
PALE_BLUE = "EAF1FA"
PALE_YELLOW = "FFF7DE"
PALE_RED = "FBEAEC"
GRAY = "667085"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=NAVY, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    paragraph.add_run(" 页")


def add_para(doc, text="", style=None, bold=False, color=None, size=None, align=None, space_after=5):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold = bold
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if size:
            r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(3)


def add_callout(doc, title, body, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    p.add_run(body)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_qa(doc, number, question, answer, evidence=None, boundary=None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, TEAL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"Q{number}. {question}")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(255, 255, 255)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("参考回答：")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run(answer)
    if evidence:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run("项目证据：")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(TEAL)
        p.add_run(evidence)
    if boundary:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run("回答边界：")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string("B54708")
        p.add_run(boundary)
    else:
        doc.paragraphs[-1].paragraph_format.space_after = Pt(8)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        shade(hdr.cells[i], TEAL)
        set_cell_text(hdr.cells[i], h, bold=True, color=WHITE)
        if widths:
            hdr.cells[i].width = Cm(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if ridx % 2 == 1:
                shade(cells[i], "F6F8FA")
            set_cell_text(cells[i], value)
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def setup_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color in [("Title", 30, NAVY), ("Heading 1", 20, NAVY), ("Heading 2", 15, TEAL), ("Heading 3", 12, NAVY)]:
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(6)

    header = sec.header.paragraphs[0]
    header.text = "APEX · AI 产品经理面试问答手册"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    header.runs[0].font.size = Pt(8.5)
    add_page_number(sec.footer.paragraphs[0])
    sec.footer.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    sec.footer.paragraphs[0].runs[0].font.size = Pt(8)
    return doc


doc = setup_doc()

# Cover
add_para(doc, "APEX", bold=True, color=TEAL, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
add_para(doc, "AI 产品经理面试问答手册", style="Title", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
add_para(doc, "乳腺筛查 AI 决策与价值评估平台", bold=True, color=NAVY, size=17, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=22)
add_callout(doc, "文档用途", "把项目中真实完成的产品设计、AI 能力、评测机制、ROI 模型、Bad Case 与协作方式，整理成可直接复习和口述的面试答案。", LIGHT)
add_para(doc, "版本：V1.0", align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY, size=10, space_after=3)
add_para(doc, "撰写人：张文琪", align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY, size=10, space_after=3)
add_para(doc, "日期：2026 年 8 月 16 日", align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY, size=10, space_after=18)
add_para(doc, "项目地址", bold=True, color=TEAL, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
add_para(doc, "GitHub：https://github.com/mxdd52077/breast-roi-copilot", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, space_after=3)
add_para(doc, "在线体验：https://breast-roi-copilot-6porxcttafmremx7ttmt88.streamlit.app", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5)
doc.add_page_break()

add_para(doc, "使用说明", style="Heading 1")
add_bullets(doc, [
    "先背“30 秒项目介绍”和“3 分钟 STAR”，再按问题分类复习。",
    "回答时先讲业务问题，再讲产品方案，最后讲技术边界；不要一上来堆 RAG、Prompt、Agent 等名词。",
    "所有演示结果都要明确为合成数据或情景模型输出，不要表述成医院真实落地收益。",
    "可以诚实说明：我主导需求定义、范围、业务规则、风险控制和验收，借助 Codex 完成工程实现、测试和调试。",
])
add_callout(doc, "一句核心原则", "这个项目不是让大模型替代医学和财务计算，而是让 AI 负责理解、检索和表达，让规则与确定性模型负责可验证计算，让人负责最终确认。", PALE_YELLOW)

add_para(doc, "目录", style="Heading 1")
add_table(doc, ["章节", "内容"], [
    ("一", "项目速览与口述模板"),
    ("二", "业务与产品设计高频问答"),
    ("三", "AI、RAG 与 Prompt 高频问答"),
    ("四", "评测、Bad Case 与安全边界"),
    ("五", "ROI、风险排序与敏感性分析"),
    ("六", "项目管理、协作与技术追问"),
    ("七", "高压追问与诚实回答"),
    ("八", "5 分钟演示脚本与复习清单"),
])
doc.add_page_break()

# Section 1
add_para(doc, "一、项目速览与口述模板", style="Heading 1")
add_para(doc, "1.1 30 秒项目介绍", style="Heading 2")
add_callout(doc, "可直接口述", "我从 0 到 1 设计并落地了一款乳腺筛查 AI 决策与价值评估平台。它先接入医院人群数据，识别筛查缺口与缺失参数，再通过 PubMed 补充循证依据，由人工确认后调用从 R Shiny 迁移的确定性 ROI 模型，最后生成管理层可读的决策报告。我的核心产品工作是拆解人机协同流程、明确 AI 与规则模型边界、设计引用和数值校验、处理 Bad Case，并把复杂专业页面简化成普通用户可完成的一站式流程。", LIGHT)

add_para(doc, "1.2 3 分钟 STAR", style="Heading 2")
add_table(doc, ["部分", "参考表达"], [
    ("S 背景", "乳腺筛查评估通常同时涉及医院人群数据、医学文献、风险排序和 ROI 计算。原流程信息分散、参数来源不透明，非技术用户很难完成。"),
    ("T 任务", "把已有 R Shiny ROI 工具升级为一套可演示、可解释、可追溯的 AI 决策产品，同时保持计算公式稳定，不让大模型直接修改 ROI。"),
    ("A 行动", "设计“数据接入—缺失项识别—循证补充—人工确认—风险排序—ROI 仿真—报告生成”流程；接入 PubMed；设计结构化 Prompt、PMID 引用校验、数值溯源、Bad Case 拦截；迁移 R 公式并编写测试；持续简化导航与页面。"),
    ("R 结果", "完成可在线演示的 Streamlit MVP，支持中英文、真实 PubMed 检索、合成数据演示、风险优先级、确定性 ROI、敏感性分析与 AI 报告；代码和测试已上传 GitHub，并部署到 Streamlit Cloud。"),
])

add_para(doc, "1.3 项目事实卡", style="Heading 2")
add_table(doc, ["维度", "当前版本"], [
    ("产品定位", "受控型 AI 决策助手 / Human-in-the-loop workflow"),
    ("目标用户", "医疗运营、筛查项目负责人、分析师、管理层"),
    ("当前范围", "40–74 岁女性，DBT / 3D mammography 筛查"),
    ("AI 负责", "自然语言理解、证据摘要、参数建议、报告起草"),
    ("规则模型负责", "Care Gap、优先级排序、ROI 计算、敏感性分析、校验"),
    ("人工负责", "确认数据、证据、参数与报告"),
    ("数据状态", "公开文献 + 合成医院数据；非真实医院实施数据"),
    ("技术栈", "Python、Streamlit、OpenAI API、NCBI PubMed API、pytest"),
])

# Section 2 Q&A
doc.add_page_break()
add_para(doc, "二、业务与产品设计高频问答", style="Heading 1")
qas = [
    (1, "你为什么做这个产品？", "因为原始 ROI 工具只能让专业人员手动填参数，无法解释参数从哪里来，也没有把医院数据、证据、风险人群和报告串起来。我把问题重新定义为“如何让非技术用户在受控流程中完成筛查决策”，而不是单纯做一个计算器。", "从 R Shiny 工具出发，最终形成一站式 AI 决策助手。", "不要声称已经在多家医院落地。"),
    (2, "目标用户是谁？核心场景是什么？", "核心用户是筛查项目运营负责人和医疗分析师，管理层是报告的最终阅读者。核心场景是：医院想把 DBT 筛查覆盖率从当前水平提升到目标水平，但缺少完整参数，需要判断应该优先联系谁、可能产生多少临床和财务价值，以及哪些假设必须先核实。", "产品分别提供数据接入、风险优先级、ROI 与管理层报告。", None),
    (3, "用户最痛的点是什么？", "不是缺少某一个模型，而是缺少完整的决策闭环：医院数据与文献分开、参数来源不清、计算过程难解释、报告缺少可行动建议。我的解决方案是把每一步的输入、来源、确认和输出都显式化。", "参数来源表、人工确认、计算过程入口和报告建议。", None),
    (4, "你的 MVP 为什么这样排序？", "P0 是数据接入、核心 ROI 和报告，因为它们构成最短价值闭环；P1 是 PubMed 循证补充和风险优先级，用来解决参数缺失和资源有限；独立 RAG 问答、复杂价值库和 AI 评测页面保留底层代码但从主导航隐藏，避免产品为了展示技术而牺牲易用性。", "项目经历了主动做减法和导航重构。", None),
    (5, "为什么后来要简化页面？", "早期版本按技术模块拆页，功能完整但用户认知成本高。通过自己的可用性走查，我发现普通用户不知道先做什么，也会把重复报告和高级统计工具混淆。因此把主路径收敛到 AI 决策助手，专业页面作为二级入口。", "隐藏重复入口、合并证据与价值库、弱化高级 ROI。", "不要虚构正式用户访谈样本量。"),
    (6, "这个产品真正的差异化是什么？", "差异化不是“有一个聊天框”，而是把非结构化证据、医院数据、确定性 ROI 和人工审批组合成可追溯的工作流。大模型的输出不能直接进入 ROI，每个关键数字都有来源或被拦截。", "PMID 白名单、数值溯源、人工确认、确定性计算。", None),
    (7, "如何定义成功指标？", "我会分三层：任务层看端到端完成率、报告生成成功率和平均完成时长；AI 质量层看引用有效率、结构化输出通过率、数值溯源通过率和 Bad Case 率；业务层看已确认参数覆盖率、优先外展带来的预期增量和方案 ROI。当前 MVP 主要验证功能可用性和规则正确性，尚无真实上线业务指标。", "已有单元测试和校验机制，但没有生产用户数据。", None),
    (8, "为什么不直接让用户自己查文献、填 Excel？", "因为决策成本不只在计算，而在把场景转成检索问题、判断证据是否可迁移、记录参数来源并保证报告数字一致。产品把这些重复工作标准化，同时保留人工审批。", None, None),
]
for q in qas:
    add_qa(doc, *q)

# Section 3
doc.add_page_break()
add_para(doc, "三、AI、RAG 与 Prompt 高频问答", style="Heading 1")
qas = [
    (9, "这个项目哪里真正用了 AI？", "AI 用在四个位置：把自然语言医院场景解析为结构化字段；对检索到的 PubMed 摘要进行循证总结；在证据允许时提出参数建议；把确定性 ROI 输出改写为管理层报告。Care Gap、风险排序、ROI 和敏感性分析不是大模型计算。", "OpenAI API 只负责语言与证据任务。", None),
    (10, "你的项目是 workflow 还是 Agent？", "当前版本更准确地说是受控工作流或有限自主 Copilot，不是开放式自主 Agent。步骤和工具顺序由产品定义，AI 在限定输入和结构中执行，关键节点必须人工确认。这更符合医疗决策的可控性要求。", "固定六步主流程和人工确认卡。", None),
    (11, "RAG 在哪里？", "Evidence Search 先通过 NCBI 官方 API 检索真实 PubMed 记录，随后 Evidence Analyst 只基于当前摘要回答，并要求每个结论引用检索结果中的 PMID。这个“先检索、后基于检索内容生成、再校验引用”的链路就是当前轻量 RAG。", "真实 PubMed API、摘要上下文、PMID 引用校验。", "当前没有向量数据库，也不应把它包装成复杂向量 RAG。"),
    (12, "为什么没有上向量数据库？", "MVP 的语料规模小、每次只处理当前检索到的少量摘要，关键词检索已经能验证核心风险：文献真实性、引用一致性和回答边界。向量库会增加切片、召回评测和运维成本，但短期不提升主要用户价值。后续在扩展指南、政策和内部文档时再引入。", None, None),
    (13, "为什么用 PubMed？其他来源呢？", "PubMed 适合临床研究与筛查效果证据，但并不覆盖所有参数。SEER 更适合发病率和分期分布，CMS 或医院财务系统更适合成本，本地医院系统更适合召回率和完成率。产品设计上先使用医院数据，只有缺失项才按参数类型去对应来源查找。", "界面明确显示参数来源类型和缺失时去哪里找。", None),
    (14, "Prompt 是怎么设计的？", "Prompt 采用五层约束：角色与受众、可用输入、禁止事项、结构化输出、失败处理。比如报告生成只允许使用已锁定 ROI 快照和已批准证据；不能新增数字或 PMID；必须区分模型结果与事实；输出固定章节，并提供可执行但不越界的下一步建议。", "报告 Prompt、证据总结 Prompt、参数提取 Prompt 分离。", None),
    (15, "你做了 Prompt 调优吗？", "做过基于 Bad Case 的迭代。早期报告会重复数字、生成无法追溯的 291、只写风险提示却没有建议，证据总结还会遗漏或新增 PMID。之后我收紧输出 schema、提供允许数字清单、加入引用白名单、要求建议与已知输入绑定，并增加一次自动纠错与失败降级。", "数值校验、PMID 校验、自动纠错、模板回退。", "自然语言和 Codex 协作本身不是产品 Prompt 调优；真正的调优是修改应用内 Prompt 并用用例验证。"),
    (16, "为什么报告还要 AI，模板不够吗？", "模板适合稳定数字展示，AI 的价值在于根据受众组织重点、解释关键驱动因素和生成有条件的行动建议。为了控制风险，数字和结论先由规则模型生成，AI 只负责叙事，失败时退回确定性模板。", None, None),
    (17, "如何防止 AI 自动改 ROI 参数？", "产品层面把建议态和确认态分开：AI 推荐只显示推荐值、区间、适用人群和证据；用户点击接受后才形成确认记录。ROI 只读取确认值，且保留原始值、推荐值和最终值。", "Parameter Copilot 和审计记录。", None),
]
for q in qas:
    add_qa(doc, *q)

# Section 4
doc.add_page_break()
add_para(doc, "四、评测、Bad Case 与安全边界", style="Heading 1")
add_para(doc, "4.1 三套评测体系不要混淆", style="Heading 2")
add_table(doc, ["对象", "评测问题", "示例指标"], [
    ("Care Gap 规则模型", "是否识别出应筛未筛人群", "敏感度、特异度、精确率、准确率、F1"),
    ("LLM / RAG", "回答是否有依据、数字是否可追溯", "PMID 有效率、引用覆盖率、结构通过率、数值溯源率、Bad Case 率"),
    ("ROI 与业务", "计算是否一致、方案是否有价值", "R/Python 对照、单元测试、净节约、ROI、敏感性、任务完成率"),
])

qas = [
    (18, "你遇到过哪些 Bad Case？", "主要有四类：AI 返回重复或不属于本次检索的 PMID；报告生成无法追溯的数字，例如 291；证据不足时仍给出确定参数；报告只强调“这是模拟”，没有给管理层可执行建议。", "这些问题都在实际开发与验收中出现过。", None),
    (19, "你如何做 Bad Case 归因？", "我按“输入、检索、生成、校验、交互”五层定位。无关文献通常是检索式和召回范围问题；假 PMID 是生成约束和引用验证问题；数字错误是上下文和数值白名单问题；报告不可用是 Prompt 目标与信息架构问题。不同原因对应不同修复，而不是统一重试。", None, None),
    (20, "Bad Case 最终怎么解决？", "引用问题通过当前检索结果 PMID 白名单拦截；数字问题通过 ROI 输出允许值列表和本地解析校验；证据不足时强制输出 insufficient evidence；报告失败先自动纠错一次，仍失败则不展示未验证文本并回退确定性模板；关键参数始终人工确认。", "校验器、自动纠错、阻断与降级。", None),
    (21, "你如何评测 RAG 检索质量？", "正式评测应建立小型黄金测试集：每个问题标注期望 PMID 或相关性等级，测 Recall@K、Precision@K、MRR，并记录无结果和错主题。当前项目已经有真实 PubMed 检索与引用校验，但还没有形成成熟的大规模检索评测看板，因此我会把它作为后续完善项。", None, "不要声称已经完成完整 RAG benchmark。"),
    (22, "模型性能页面的敏感度、特异度从哪里来？", "它评估的是合成数据中的 Care Gap 识别，不是大模型。系统用 care_gap_score 与阈值生成预测，再与合成 ground_truth_gap 比较形成混淆矩阵。由于真值也是按规则生成，这只能证明演示框架和计算正确，不能代表真实医院模型表现。", "合成 10,000 人数据和阈值页面。", None),
    (23, "为什么不能让大模型直接判断 Care Gap？", "Care Gap 主要由筛查间隔、最近筛查日期和预约状态等结构化规则决定，规则模型更稳定、便宜、可审计。大模型可用于解释非结构化备注，但不应该替代清晰的业务规则。", None, None),
    (24, "医疗场景如何控制幻觉和合规风险？", "坚持最小权限与能力隔离：不上传密钥、不把患者数据持久化、演示数据明确标注合成；LLM 不计算 ROI、不新增 PMID、不自动改参数；证据和数字均本地校验；无充分依据时显示不足，而不是强行回答。", "secrets.toml 被 gitignore；部署时通过 Secrets 配置。", None),
]
for q in qas:
    add_qa(doc, *q)

# Section 5
doc.add_page_break()
add_para(doc, "五、ROI、风险排序与敏感性分析", style="Heading 1")
qas = [
    (25, "原来的 R 模型在哪里用上了？", "核心 ROI 公式被迁移到 Python 计算模块，Streamlit 页面只是收集参数和展示结果。新增筛查人数、检出病例、生命挽救、筛查与随访成本、阶段转移节约、净节约和 ROI 都沿用 R 模型逻辑，并通过单元测试和示例数值对照验证。", "计算代码与页面分离，LLM 不参与数学计算。", None),
    (26, "ROI 是怎么计算的？", "先计算新增筛查人数，再根据检出率和年龄校正得到新增病例；用分期分布、分期转移假设和分期成本差估算避免治疗成本；项目成本包含筛查与召回随访；净节约等于避免治疗成本减项目成本，ROI 等于净节约除以项目成本。", "README 中保留公式地图。", "这些是情景模型输出，不是已实现收益。"),
    (27, "风险优先级页面做什么？", "它把所有符合条件的人计算优先级，再从中截取可用外展名额。例如 10,000 人都参与排序，但如果只有 2,000 个名额，页面重点展示排名前 2,000 人，并与随机外展比较预期覆盖、完成和检出价值。", "透明权重：60% Care Gap、25% 距上次筛查时间、15% 完成概率。", "该排序逻辑是产品新增能力，不是原 R 模型自带。"),
    (28, "为什么这样设置 60/25/15？", "这是可解释的 MVP 业务权重：先保证真正存在筛查缺口的人优先，其次考虑逾期程度，最后兼顾外展可转化性。它不是训练出的临床最优权重，因此页面应标注为演示规则；上线前需用真实数据、业务目标和公平性评估重新校准。", None, None),
    (29, "什么是敏感性分析？", "它回答“哪个不确定参数最影响结果”。系统一次只把一个关键参数上下调整 20%，其余不变，然后重新跑同一套 ROI 公式，观察净节约变化。波动越大，说明决策越依赖这个参数，管理层越应优先核实。", "当前显示分期改善、检出率、筛查成本、随访成本和召回率。", None),
    (30, "什么是分期改善假设？", "它假设更早筛查会让部分原本在区域期或远处期发现的病例提前到更早分期，从而减少治疗成本。它对 ROI 影响很大，但当前批准证据没有给出可直接迁移的比例，所以必须把它当作高风险假设，而不是已证实事实。", "敏感性分析将其识别为关键驱动因素。", None),
    (31, "保守、基准、积极三种情景怎么来的？", "基准使用当前确认值；保守情景把检出与分期改善等效果参数下调、把筛查和随访成本上调；积极情景反向调整。三组参数都进入同一套确定性公式。这是透明压力测试，不是在预测未来发生概率，也不是置信区间。", "页面提供计算说明脚注。", None),
    (32, "医院上传数据后，44.71% 和平均年龄 57 是怎么来的？", "目标人群规模来自上传记录行数，平均年龄是 age 字段均值。当前筛查率由系统根据基准日、最近筛查日期、是否从未筛查、预约状态和筛查间隔规则派生，再统计当前处于筛查覆盖状态的人数占比。目标筛查率则由用户输入。", "参数来源表区分医院数据、医院输入、缺失待补充。", None),
]
for q in qas:
    add_qa(doc, *q)

# Section 6
doc.add_page_break()
add_para(doc, "六、项目管理、协作与技术追问", style="Heading 1")
qas = [
    (33, "你作为 AI 产品经理具体做了什么？", "我负责定义目标用户和决策链路、拆解输入输出、划定模型边界、设计人机协同和异常兜底、制定验收规则、复盘 Bad Case，并根据可用性持续做功能取舍。工程上我借助 Codex 完成代码迁移、模块实现、测试和调试，但每个业务口径与验收决策由我确认。", "PRD、流程、Prompt 规则、测试用例和多轮产品重构。", None),
    (34, "你和 Codex 如何协作？", "我把需求拆成可验收的小步骤，例如先迁移 R 公式，再做真实 PubMed 搜索，再做有引用的问答，最后才接参数推荐和报告。Codex 负责读取代码、实现与测试；我通过页面结果、数值对照和风险规则验收，并在发现问题时调整需求。这个过程类似 PM 与研发的快速共创。", None, None),
    (35, "你如何保证不是只靠 Vibe Coding？", "我没有用“一句话生成完整产品”的方式，而是保留模块边界、单元测试、输入 schema、数值对照和 Git 版本。对关键功能先定义不可违反的验收标准，例如不得虚构 PMID、LLM 不计算 ROI、参数必须人工确认，再实现和回归。", "pytest、模块化目录、GitHub 提交、Streamlit 部署。", None),
    (36, "最大的技术挑战是什么？", "最大的挑战不是调 API，而是不同模型的边界：医学证据可能不足，ROI 又需要确定参数。如果直接让 LLM 填数会放大风险。我把“建议—确认—计算”拆开，并通过引用和数值验证让失败显式化。", None, None),
    (37, "最大的产品挑战是什么？", "早期功能过多、导航复杂，用户看不懂各页面关系。我最终把主路径收敛到 AI 决策助手，隐藏重复入口，把高级仿真和证据治理降级为专业工具。这体现了从“展示能力”转向“完成任务”的产品取舍。", None, None),
    (38, "为什么选择 Streamlit？", "它适合快速验证数据产品和 AI 工作流，可以低成本完成参数控件、表格、图表、上传和多页面部署。MVP 的目标是验证流程和规则，不是构建高并发生产系统；验证后再考虑前后端分离。", None, None),
    (39, "如何部署和管理密钥？", "代码托管在 GitHub，应用部署到 Streamlit Community Cloud。OpenAI Key 不进入仓库，本地放在被 gitignore 的 .streamlit/secrets.toml，云端通过 Streamlit Secrets 配置。", "GitHub 与在线体验地址可现场展示。", None),
    (40, "下一版你会做什么？", "第一，建立 20–30 条高质量黄金测试集，量化引用、数字和建议质量；第二，接入 SEER/CMS 或可配置外部数据源，按参数类型选来源；第三，用真实脱敏数据校准 Care Gap 和风险排序；第四，增加报告版本对比、成本和延迟监控。", None, None),
]
for q in qas:
    add_qa(doc, *q)

# Section 7
doc.add_page_break()
add_para(doc, "七、高压追问与诚实回答", style="Heading 1")
add_table(doc, ["追问", "建议回答"], [
    ("这不就是一个规则计算器吗？", "ROI 本身确实是确定性计算，这是刻意设计。AI 解决的是非结构化证据、参数缺失和报告表达；规则模型保证核心决策可验证。产品价值在双引擎协同，不在强行把所有功能 AI 化。"),
    ("没有向量库也叫 RAG？", "RAG 的核心是先检索再基于检索内容生成。当前小规模摘要使用 PubMed API + 上下文注入 + 引用校验已经满足 MVP；向量库是规模化技术选项，不是定义本身。"),
    ("为什么不是 Agent？", "医疗场景更需要受控步骤和人工确认。当前定位是有限自主 Copilot；如果未来引入工具选择和循环规划，也会保留权限、预算和审批边界。"),
    ("这些 ROI 数字可信吗？", "公式经过 R/Python 对照和单元测试，说明计算一致；但输入参数和真实实施效果仍需本地验证，所以只能称情景估算，不能称已实现收益。"),
    ("你的模型评测为什么这么高？", "那是合成规则数据上的 Care Gap 识别性能，只验证流程，不代表真实医院泛化性能。我会主动说明真值也是规则生成的。"),
    ("你真的做过 Prompt 调优吗？", "做过针对应用 Prompt 的 Bad Case 迭代：引用越界、数字幻觉、证据不足和建议缺失；通过 schema、白名单、校验和回退优化。不是把我和 Codex 的聊天泛称为 Prompt 调优。"),
    ("项目产生了什么业务结果？", "当前完成的是可在线演示、可测试的 MVP，验证了端到端流程和风险控制；尚无真实医院上线数据，因此不虚构降本金额、用户增长或效率提升百分比。"),
    ("为什么只支持 DBT？", "为了控制模型边界。不同筛查方式需要独立检出率、召回率、成本和适用人群，不应把 MRI、超声或数字乳腺摄影混用同一参数。"),
    ("如果证据和医院参数冲突怎么办？", "不自动覆盖。页面同时展示本地值、外部建议、适用人群和证据强度，由业务与临床负责人确认最终值，并保留修改记录。"),
    ("如果 API 失败怎么办？", "检索失败显示明确错误，不生成虚构文献；LLM 失败则使用演示文章或确定性报告模板，ROI 主功能仍可运行。"),
])

add_para(doc, "不应使用的表述", style="Heading 2")
add_table(doc, ["不要说", "建议改成"], [
    ("已落地多家医院", "已完成可在线演示的 MVP"),
    ("模型误差降低 92%", "通过单元测试和 R/Python 数值对照验证公式一致性"),
    ("AI 自动推荐最优参数", "AI 提供有证据边界的参数候选，最终人工确认"),
    ("Agent 自动完成医疗决策", "受控工作流辅助决策，关键节点人工审批"),
    ("RAG 保证医学正确", "RAG 提高可追溯性，结论仍受检索范围和证据质量限制"),
    ("合成数据评测证明模型有效", "合成数据用于验证流程和指标计算"),
])

# Section 8
doc.add_page_break()
add_para(doc, "八、5 分钟演示脚本与复习清单", style="Heading 1")
add_para(doc, "8.1 5 分钟现场演示", style="Heading 2")
add_table(doc, ["时间", "操作", "讲解重点"], [
    ("0:00–0:40", "打开 AI 决策助手", "一句话介绍目标用户、业务问题和双引擎边界。"),
    ("0:40–1:30", "上传合成医院 CSV", "展示当前筛查率、平均年龄如何从医院字段派生，并说明数据为合成。"),
    ("1:30–2:20", "查看缺失参数与证据补充", "说明只对缺失项查对应来源，PubMed 不是所有参数的唯一来源。"),
    ("2:20–3:00", "确认参数并运行 ROI", "强调 LLM 不计算 ROI，结果来自 R 公式迁移的确定性模型。"),
    ("3:00–3:45", "展示风险优先级", "所有人先排序，再按有限外展名额截取 Top N；权重透明可解释。"),
    ("3:45–4:30", "展示敏感性与三情景", "说明不是预测，而是 ±20% 压力测试，用来找最该核实的参数。"),
    ("4:30–5:00", "生成管理层报告", "展示 AI 只组织叙事和建议，数字经过溯源校验，失败可降级。"),
])

add_para(doc, "8.2 面试前最后检查", style="Heading 2")
add_bullets(doc, [
    "能在 30 秒内说清：用户、问题、方案、AI 边界、当前结果。",
    "能画出主流程，并解释每一步为什么存在。",
    "能说出至少 3 个真实 Bad Case、原因和修复。",
    "能区分 Care Gap 评测、LLM/RAG 评测和 ROI 评测。",
    "能解释敏感性分析和三情景不是预测。",
    "能主动说明合成数据、证据不足与未真实落地的边界。",
    "提前打开在线网址，并准备本地或截图备用。",
])

add_para(doc, "8.3 一页速记", style="Heading 2")
add_callout(doc, "项目主线", "医院数据 → 缺失项 → 外部证据 → 人工确认 → 风险排序 → 确定性 ROI → AI 报告", LIGHT)
add_callout(doc, "三层能力", "AI 负责理解和表达；规则模型负责计算和校验；人工负责最终决策。", PALE_BLUE)
add_callout(doc, "四类 Bad Case", "PMID 越界、数字不可追溯、证据不足仍下结论、报告无行动建议。", PALE_RED)
add_callout(doc, "一句边界", "这是基于合成数据和已确认假设的决策支持 MVP，不是医疗建议，也不代表已实现的临床或财务结果。", PALE_YELLOW)

add_para(doc, "附录：常用名词", style="Heading 1")
add_table(doc, ["名词", "小白解释"], [
    ("Care Gap", "按照筛查规则本该筛查、但目前没有完成或已经逾期的人群缺口。"),
    ("RAG", "先找资料，再让模型只根据找到的资料回答。"),
    ("PMID", "每篇 PubMed 文献的唯一编号，可用来核对引用是否真实。"),
    ("Prompt", "给大模型的任务说明、输入边界、输出格式和禁止事项。"),
    ("Bad Case", "模型输出不符合预期或存在风险的失败案例。"),
    ("Human-in-the-loop", "AI 给建议，但关键结果由人确认后才能生效。"),
    ("敏感性分析", "改变一个假设并重新计算，看结果对哪个参数最敏感。"),
    ("Stage shift", "假设部分癌症因为筛查而在更早分期被发现。"),
    ("确定性模型", "同样输入一定得到同样输出，不依赖大模型随机生成。"),
])

add_para(doc, "参考链接", style="Heading 2")
add_bullets(doc, [
    "GitHub：https://github.com/mxdd52077/breast-roi-copilot",
    "在线应用：https://breast-roi-copilot-6porxcttafmremx7ttmt88.streamlit.app",
    "项目 README：GitHub 仓库首页",
])

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(str(OUTPUT))
