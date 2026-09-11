from datetime import date
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from src.patient_companion.documents import DocumentParseError, parse_document
from src.patient_companion.demo_provider import build_demo_case, get_demo_case_summaries
from src.patient_companion.rag import answer_from_demo, retrieve_chunks
from src.patient_companion.provider import _safe_openai_error, assemble_patient_case
from src.patient_companion.review import apply_fact_review
from src.patient_companion.schemas import (
    CareEvent,
    ClinicalFact,
    ConfirmationStatus,
    EventType,
    KnowledgeChunk,
    PatientCase,
    PathologyMarker,
    PatientProfile,
    PatientExtraction,
    SourceDocument,
    resolve_confirmation_status,
)
from src.patient_companion.tasks import build_care_tasks
from src.patient_companion.validation import ProvenanceError, validate_case_provenance


def test_txt_document_is_split_into_traceable_pages() -> None:
    document = parse_document(
        filename="医嘱.txt",
        content="第一段医嘱\f第二段复诊安排".encode("utf-8"),
    )

    assert document.filename == "医嘱.txt"
    assert [page.page_number for page in document.pages] == [1, 2]
    assert document.pages[1].text == "第二段复诊安排"


def test_empty_or_unsupported_documents_are_rejected() -> None:
    with pytest.raises(DocumentParseError, match="没有可读取的文字"):
        parse_document(filename="empty.txt", content=b"   ")

    with pytest.raises(DocumentParseError, match="仅支持"):
        parse_document(filename="scan.png", content=b"not-an-image")


def test_demo_library_contains_distinct_traceable_care_stages() -> None:
    summaries = get_demo_case_summaries()
    cases = [build_demo_case(summary.id) for summary in summaries]

    assert [summary.id for summary in summaries] == [
        "endocrine",
        "chemotherapy",
        "radiotherapy",
    ]
    assert len({case.profile.treatment_stage for case in cases}) == 3
    for case in cases:
        validate_case_provenance(case.documents, case.facts, case.events)
        assert case.facts
        assert case.events


def test_unknown_demo_case_is_rejected() -> None:
    with pytest.raises(ValueError, match="未知的模拟档案"):
        build_demo_case("not-a-case")


@pytest.mark.parametrize(
    ("error_name", "status_code", "expected"),
    [
        ("AuthenticationError", 401, "Key 无效"),
        ("RateLimitError", 429, "额度不足"),
        ("APITimeoutError", None, "连接 OpenAI 超时"),
        ("BadRequestError", 400, "结构化请求"),
    ],
)
def test_openai_failures_are_mapped_to_actionable_safe_messages(
    error_name: str, status_code: int | None, expected: str
) -> None:
    error_type = type(error_name, (Exception,), {})
    error = error_type("sensitive upstream details")
    error.status_code = status_code

    message = _safe_openai_error(error, "患者资料分析")

    assert expected in message
    assert "sensitive upstream details" not in message


def test_docx_document_text_is_extracted() -> None:
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("复诊安排：下周一乳腺外科复诊。")
    document.save(buffer)

    parsed = parse_document(filename="复诊.docx", content=buffer.getvalue())
    assert parsed.pages[0].text == "复诊安排：下周一乳腺外科复诊。"


def test_image_only_pdf_is_rejected_as_no_readable_text() -> None:
    buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(buffer)

    with pytest.raises(DocumentParseError, match="扫描件和图片暂不支持"):
        parse_document(filename="扫描件.pdf", content=buffer.getvalue())


def test_confirmed_medical_events_become_traceable_tasks() -> None:
    event = CareEvent(
        id="event-1",
        event_type=EventType.MEDICATION,
        title="服用来曲唑 2.5 mg",
        scheduled_date=date(2026, 8, 23),
        frequency="每日一次",
        source_fact_ids=["fact-1"],
        confirmation_status=ConfirmationStatus.CONFIRMED,
    )

    tasks = build_care_tasks([event], today=date(2026, 8, 22))

    assert len(tasks) == 1
    assert tasks[0].source_event_id == "event-1"
    assert tasks[0].due_date == date(2026, 8, 23)
    assert tasks[0].instructions == "服用来曲唑 2.5 mg"


def test_unconfirmed_medical_events_never_become_actionable_tasks() -> None:
    event = CareEvent(
        id="event-2",
        event_type=EventType.MEDICATION,
        title="疑似用药信息",
        scheduled_date=date(2026, 8, 23),
        source_fact_ids=["fact-2"],
        confirmation_status=ConfirmationStatus.NEEDS_CONFIRMATION,
    )

    assert build_care_tasks([event], today=date(2026, 8, 22)) == []


def test_empty_confirmation_widget_value_preserves_the_current_status() -> None:
    current = ConfirmationStatus.CONFIRMED

    assert resolve_confirmation_status(None, current) is current
    assert resolve_confirmation_status("已排除", current) is ConfirmationStatus.REJECTED


def test_rag_returns_only_known_sources_and_refuses_unsupported_questions() -> None:
    chunks = [
        KnowledgeChunk(
            id="nci-fatigue",
            title="癌症治疗副作用",
            publisher="NCI",
            source_url="https://www.cancer.gov/about-cancer/treatment/side-effects",
            topic="副作用",
            treatment_stages=["治疗期"],
            jurisdiction="国际",
            updated_at=date(2026, 1, 1),
            reviewed_at=date(2026, 8, 1),
            risk_level="一般",
            text="疲劳是癌症治疗期间常见的副作用。请把症状告诉医疗团队。",
        )
    ]

    retrieved = retrieve_chunks("治疗期间疲劳怎么办", chunks, stage="治疗期")
    answer = answer_from_demo("治疗期间疲劳怎么办", retrieved)

    assert answer.status == "supported"
    assert answer.citations == ["nci-fatigue"]
    assert "医疗团队" in answer.answer

    unsupported = answer_from_demo("我应该把药量加倍吗", [])
    assert unsupported.status == "insufficient_evidence"
    assert unsupported.citations == []


def test_clinical_facts_require_source_provenance() -> None:
    with pytest.raises(ValueError):
        ClinicalFact(
            id="fact-no-source",
            category="medication",
            value="来曲唑",
            original_text="来曲唑",
            source_document_id="",
            source_page=1,
        )


def test_patient_profile_is_independent_from_b2b_roi_models() -> None:
    profile = PatientProfile(
        display_name="演示患者",
        treatment_stage="治疗期",
        diagnosis_summary="乳腺癌术后辅助治疗",
    )

    assert profile.display_name == "演示患者"
    assert "population_size" not in PatientProfile.model_fields


def test_patient_profile_uses_strict_pathology_marker_pairs() -> None:
    profile = PatientProfile(
        display_name="演示患者",
        treatment_stage="治疗期",
        diagnosis_summary="乳腺癌术后",
        pathology_markers=[PathologyMarker(name="ER", value="阳性 90%")],
    )

    assert profile.pathology_markers[0].name == "ER"


def test_case_validation_rejects_facts_not_present_in_source() -> None:
    document = SourceDocument(
        id="doc-1",
        filename="医嘱.txt",
        media_type="text/plain",
        uploaded_at="2026-08-22T00:00:00Z",
        pages=[{"page_number": 1, "text": "医嘱：两周后复诊。"}],
    )
    invented = ClinicalFact(
        id="fact-invented",
        category="用药",
        value="增加药量",
        original_text="增加药量",
        source_document_id="doc-1",
        source_page=1,
    )

    with pytest.raises(ProvenanceError, match="无法在原文中定位"):
        validate_case_provenance([document], [invented], [])


def test_case_validation_accepts_pdf_line_breaks_inside_verbatim_fact() -> None:
    document = SourceDocument(
        id="doc-1",
        filename="医嘱.pdf",
        media_type="application/pdf",
        uploaded_at="2026-08-22T00:00:00Z",
        pages=[
            {
                "page_number": 1,
                "text": "用药医嘱：来曲唑 2.5 mg，\n每日一次。",
            }
        ],
    )
    fact = ClinicalFact(
        id="fact-1",
        category="用药医嘱",
        value="来曲唑 2.5 mg，每日一次。",
        original_text="来曲唑 2.5 mg，每日一次。",
        source_document_id="doc-1",
        source_page=1,
    )

    validate_case_provenance([document], [fact], [])


def test_case_validation_rejects_event_with_unknown_fact() -> None:
    event = CareEvent(
        id="event-invented",
        event_type=EventType.TREATMENT,
        title="新增治疗",
        scheduled_date=date(2026, 8, 30),
        source_fact_ids=["missing-fact"],
    )

    with pytest.raises(ProvenanceError, match="不存在的事实"):
        validate_case_provenance([], [], [event])


def test_provider_assembly_keeps_uploaded_documents_outside_model_output() -> None:
    document = SourceDocument(
        id="doc-1",
        filename="病历.txt",
        media_type="text/plain",
        uploaded_at="2026-08-22T00:00:00Z",
        pages=[{"page_number": 1, "text": "诊断：乳腺癌术后。"}],
    )
    extraction = PatientExtraction(
        profile=PatientProfile(
            display_name="患者",
            treatment_stage="治疗期",
            diagnosis_summary="乳腺癌术后",
        ),
        facts=[
            ClinicalFact(
                id="fact-1",
                category="诊断",
                value="乳腺癌术后",
                original_text="乳腺癌术后",
                source_document_id="doc-1",
                source_page=1,
            )
        ],
        events=[],
    )

    case = assemble_patient_case([document], extraction)
    assert case.documents == [document]
    assert case.tasks == []


def test_fact_review_atomically_updates_linked_event_status() -> None:
    case = PatientCase(
        profile=PatientProfile(
            display_name="患者",
            treatment_stage="治疗期",
            diagnosis_summary="乳腺癌术后",
        ),
        facts=[
            ClinicalFact(
                id="fact-1",
                category="复诊",
                value="两周后复诊",
                original_text="两周后复诊",
                source_document_id="doc-1",
                source_page=1,
            )
        ],
        events=[
            CareEvent(
                id="event-1",
                event_type=EventType.APPOINTMENT,
                title="乳腺外科复诊",
                scheduled_date=date(2026, 9, 10),
                source_fact_ids=["fact-1"],
            )
        ],
    )

    confirmed = apply_fact_review(case, {"fact-1": "已确认"})
    rejected = apply_fact_review(case, {"fact-1": "已排除"})

    assert confirmed.events[0].confirmation_status is ConfirmationStatus.CONFIRMED
    assert rejected.events[0].confirmation_status is ConfirmationStatus.REJECTED
    assert case.facts[0].confirmation_status is ConfirmationStatus.NEEDS_CONFIRMATION


def test_fact_review_rejects_unknown_fact_without_partial_mutation() -> None:
    case = PatientCase(
        profile=PatientProfile(
            display_name="患者",
            treatment_stage="治疗期",
            diagnosis_summary="乳腺癌术后",
        ),
        facts=[],
    )

    with pytest.raises(ValueError, match="未知事实"):
        apply_fact_review(case, {"invented": "已确认"})
